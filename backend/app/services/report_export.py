"""검토 결과 스냅샷을 PDF·Excel·근거 Manifest 패키지로 내보내는 서비스."""

from __future__ import annotations

import csv
import hashlib
import html
import json
import textwrap
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from sqlalchemy import func, select

from ..config import get_settings
from ..database import ApprovalHistory, Building, DrawingChangeCandidate, EvidenceReference, ProcurementPriceResult, ReviewExport, ReviewWarning, SessionLocal, SourceFile, StandardizedItem, WorkPackage

# 원본 자료가 반복 적재된 경우에도 사용자가 수십 초 이상 기다리지 않도록
# 상세 파일은 최신순으로 제한한다. 전체 건수와 생략 여부는 Manifest에 기록한다.
EXPORT_DETAIL_LIMIT = 20_000


def _json(value: object) -> str:
    return json.dumps(value, ensure_ascii=False, default=str)


def _write_csv(path: Path, fieldnames: list[str], rows: list[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def _limited(rows: list[Any], limit: int = EXPORT_DETAIL_LIMIT) -> tuple[list[Any], bool]:
    """내보내기 상세 행을 제한하고 잘림 여부를 반환한다."""
    return rows[:limit], len(rows) > limit


def _build_xlsx(path: Path, sheets: dict[str, tuple[list[str], list[dict[str, Any]]]]) -> str | None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        return None
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name, (headers, rows) in sheets.items():
        sheet = workbook.create_sheet(sheet_name[:31])
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        for row in rows:
            sheet.append([row.get(header) for header in headers])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            width = min(60, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
            sheet.column_dimensions[column[0].column_letter].width = width
    workbook.save(path)
    return str(path)


def _build_pdf(path: Path, project_id: str, counts: dict[str, int], pending: int, filters: dict[str, Any], warning_rows: list[dict[str, Any]], drawing_rows: list[dict[str, Any]], price_rows: list[dict[str, Any]]) -> str:
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.pdfgen import canvas
    except ImportError:
        html_path = path.with_suffix(".html")
        html_path.write_text(f"<html><meta charset='utf-8'><h1>공사비 적정성 검토 요약</h1><p>프로젝트: {html.escape(project_id)}</p><p>생성일: {datetime.now(timezone.utc).isoformat()}</p><p>검토 대기: {pending}건</p><pre>{html.escape(_json(counts))}</pre></html>", encoding="utf-8")
        return str(html_path)
    font_name = "Helvetica"
    for candidate in (Path(r"C:\Windows\Fonts\malgun.ttf"), Path(r"C:\Windows\Fonts\malgunsl.ttf"), Path("/usr/share/fonts/truetype/nanum/NanumGothic.ttf"), Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc")):
        if candidate.exists():
            try:
                pdfmetrics.registerFont(TTFont("Korean", str(candidate)))
                font_name = "Korean"
                break
            except Exception:
                continue
    pdf = canvas.Canvas(str(path), pagesize=A4)
    pdf.setTitle("공사비 적정성 검토 요약")
    pdf.setFont(font_name, 16)
    pdf.drawString(48, 800, "공사비 적정성 검토 결과")
    pdf.setFont(font_name, 9)
    pdf.drawString(48, 780, f"프로젝트: {project_id}")
    pdf.drawString(48, 765, f"생성일: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M')}")
    y = 735
    pdf.setFont(font_name, 10)
    pdf.drawString(48, y, "적용 조건")
    y -= 16
    filter_labels = {"review_stage": "검토 단계", "building_id": "건물", "work_package_id": "공종", "status": "상태", "warning_severity": "심각도"}
    for key, value in filters.items():
        if value:
            pdf.drawString(60, y, f"{filter_labels.get(key, key)}: {value}")
            y -= 15
    y -= 8
    pdf.setFont(font_name, 11)
    for label, value in [("경고", counts["warnings"]), ("도면 변경 후보", counts["drawings"]), ("단가 후보", counts["prices"]), ("승인 이력", counts["approvals"]), ("검토 대기", pending)]:
        pdf.drawString(60, y, f"{label}: {value}건")
        y -= 18
    y -= 8
    pdf.setFont(font_name, 9)
    pdf.drawString(48, y, "주요 검토 결과")
    y -= 15
    detail_lines = [("경고", row.get("title") or row.get("warning_type") or "검토 항목") for row in warning_rows[:12]]
    detail_lines += [("도면", f"{row.get('drawing_number') or row.get('sheet_number') or '-'} · {row.get('change_type') or '-'}") for row in drawing_rows[:8]]
    detail_lines += [("단가", f"{row.get('item_name') or '-'} · {row.get('lookup_status') or '-'}") for row in price_rows[:8]]
    if not detail_lines:
        detail_lines = [("안내", "선택한 조건에 해당하는 결과가 없습니다.")]
    for category, text in detail_lines:
        clean_text = " ".join(str(text).split())
        for line in textwrap.wrap(f"[{category}] {clean_text}", width=72, break_long_words=False, break_on_hyphens=False) or [f"[{category}] -"]:
            if y < 55:
                pdf.showPage(); pdf.setFont(font_name, 9); y = 800
            pdf.drawString(60, y, line); y -= 13
    if y < 55:
        pdf.showPage(); y = 800
    pdf.setFont(font_name, 8)
    pdf.drawString(48, y - 8, "수량·금액·단가는 승인 전 검토 후보이며 자동 확정되지 않습니다.")
    pdf.save()
    return str(path)


def execute_review_export(export_id: str) -> None:
    db = SessionLocal()
    export = db.get(ReviewExport, export_id)
    try:
        if not export:
            return
        export.status = "running"
        db.commit()
        project_dir = get_settings().export_dir / export.project_id / export.id
        project_dir.mkdir(parents=True, exist_ok=True)
        filters = json.loads(export.filter_snapshot or "{}")
        warning_query = select(ReviewWarning).where(ReviewWarning.project_id == export.project_id)
        if filters.get("review_stage") == "initial":
            warning_query = warning_query.where(ReviewWarning.warning_type.in_(["quantity_estimate_link", "formula_quantity_missing", "field_mismatch", "mapping_mismatch"]))
        elif filters.get("review_stage") == "change":
            warning_query = warning_query.where(ReviewWarning.drawing_change_candidate_id.is_not(None))
        elif filters.get("review_stage") == "price":
            warning_query = warning_query.where(ReviewWarning.warning_type.like("%price%"))
        if filters.get("building_id") or filters.get("work_package_id"):
            warning_query = warning_query.join(StandardizedItem, ReviewWarning.standardized_item_id == StandardizedItem.id, isouter=True)
            if filters.get("building_id"):
                warning_query = warning_query.where(StandardizedItem.building_id == filters["building_id"])
            if filters.get("work_package_id"):
                warning_query = warning_query.where(StandardizedItem.work_package_id == filters["work_package_id"])
        if filters.get("warning_severity"):
            warning_query = warning_query.where(ReviewWarning.severity == filters["warning_severity"])
        if filters.get("status"):
            warning_query = warning_query.where(ReviewWarning.status == filters["status"])
        # limit+1 행만 읽어 대용량 원본을 메모리와 Excel에 그대로 복사하지 않는다.
        filtered_warning_total = db.scalar(select(func.count()).select_from(warning_query.subquery())) or 0
        filtered_pending_total = db.scalar(select(func.count()).select_from(warning_query.where(ReviewWarning.status.not_in({"승인", "확정"})).subquery())) or 0
        warnings, warnings_truncated = _limited(db.scalars(warning_query.order_by(ReviewWarning.created_at.desc()).limit(EXPORT_DETAIL_LIMIT + 1)).all())
        drawing_query = select(DrawingChangeCandidate).where(DrawingChangeCandidate.project_id == export.project_id)
        if filters.get("review_stage") and filters.get("review_stage") != "change":
            drawing_query = drawing_query.where(False)
        drawings, drawings_truncated = _limited(db.scalars(drawing_query.order_by(DrawingChangeCandidate.created_at.desc()).limit(EXPORT_DETAIL_LIMIT + 1)).all())
        price_query = select(ProcurementPriceResult).where(ProcurementPriceResult.project_id == export.project_id)
        if filters.get("review_stage") and filters.get("review_stage") != "price":
            price_query = price_query.where(False)
        if filters.get("work_package_id"):
            price_query = price_query.where(ProcurementPriceResult.work_package == filters["work_package_id"])
        prices, prices_truncated = _limited(db.scalars(price_query.order_by(ProcurementPriceResult.created_at.desc()).limit(EXPORT_DETAIL_LIMIT + 1)).all())
        approvals, approvals_truncated = _limited(db.scalars(select(ApprovalHistory).where(ApprovalHistory.project_id == export.project_id).order_by(ApprovalHistory.created_at.desc()).limit(EXPORT_DETAIL_LIMIT + 1)).all())
        evidence, evidence_truncated = _limited(db.scalars(select(EvidenceReference).where(EvidenceReference.project_id == export.project_id).order_by(EvidenceReference.created_at.desc()).limit(EXPORT_DETAIL_LIMIT + 1)).all())
        files, files_truncated = _limited(db.scalars(select(SourceFile).where(SourceFile.project_id == export.project_id).order_by(SourceFile.created_at.desc()).limit(EXPORT_DETAIL_LIMIT + 1)).all())

        warning_rows = [{"id": item.id, "warning_type": item.warning_type, "severity": item.severity, "title": item.title, "detail": item.detail, "status": item.status, "rule_code": item.rule_code, "created_at": item.created_at.isoformat() if item.created_at else ""} for item in warnings]
        drawing_rows = [{"id": item.id, "drawing_number": item.drawing_number, "change_type": item.change_type, "location_ref": item.location_ref, "confidence": item.confidence, "status": item.status, "baseline_revision": item.baseline_revision, "changed_revision": item.changed_revision, "sheet_number": item.sheet_number} for item in drawings]
        price_rows = [{"id": item.id, "candidate_id": item.candidate_id, "item_name": item.item_name, "specification": item.specification, "unit": item.unit, "price": float(item.price) if item.price is not None else None, "lookup_status": item.lookup_status, "service_name": item.service_name, "source_file_id": item.source_file_id, "reference_date": item.reference_date.isoformat() if item.reference_date else ""} for item in prices]
        approval_rows = [{"id": item.id, "source_id": item.source_id, "department": item.department, "decision": item.decision, "reviewer": item.reviewer, "comment": item.comment, "evidence_ref": item.evidence_ref, "created_at": item.created_at.isoformat() if item.created_at else ""} for item in approvals]
        evidence_rows = [{"id": item.id, "warning_id": item.warning_id, "source_file_id": item.source_file_id, "file_path": item.file_path, "sheet_name": item.sheet_name, "row_ref": item.row_ref, "cell_ref": item.cell_ref, "page_number": item.page_number, "location_text": item.location_text, "extraction_confidence": item.extraction_confidence, "evidence_note": item.evidence_note} for item in evidence]
        file_rows = [{"id": item.id, "original_name": item.original_name, "file_path": item.file_path, "sha256": item.sha256, "document_type": item.document_type, "revision": item.revision, "drawing_number": item.drawing_number, "sheet_name": item.sheet_name, "is_valid": item.is_valid} for item in files]
        counts = {"warnings": len(warning_rows), "drawings": len(drawing_rows), "prices": len(price_rows), "approvals": len(approval_rows), "evidence": len(evidence_rows)}
        truncated_sections = [name for name, truncated in [("warnings", warnings_truncated), ("drawings", drawings_truncated), ("prices", prices_truncated), ("approvals", approvals_truncated), ("evidence", evidence_truncated), ("source_files", files_truncated)] if truncated]
        # 요약의 전체 건수도 상세 파일과 동일한 단계/공종 조건을 사용한다.
        # (승인 이력·근거·원본은 프로젝트 스냅샷 범위이므로 별도 집계한다.)
        filtered_drawing_total = db.scalar(select(func.count()).select_from(drawing_query.subquery())) or 0
        filtered_price_total = db.scalar(select(func.count()).select_from(price_query.subquery())) or 0
        total_counts = {"warnings": filtered_warning_total or len(warning_rows), "drawings": filtered_drawing_total or len(drawing_rows), "prices": filtered_price_total or len(price_rows), "approvals": db.scalar(select(func.count()).select_from(ApprovalHistory).where(ApprovalHistory.project_id == export.project_id)) or len(approval_rows), "evidence": db.scalar(select(func.count()).select_from(EvidenceReference).where(EvidenceReference.project_id == export.project_id)) or len(evidence_rows), "source_files": db.scalar(select(func.count()).select_from(SourceFile).where(SourceFile.project_id == export.project_id)) or len(file_rows)}
        pending = filtered_pending_total or sum(1 for item in warning_rows if item["status"] not in {"승인", "확정"})
        manifest = {"export_id": export.id, "project_id": export.project_id, "review_run_id": export.review_run_id, "filter_snapshot": filters, "approval_snapshot": json.loads(export.approval_snapshot or "{}"), "rule_version": export.rule_version, "input_hash": export.input_hash, "data_as_of": export.data_as_of.isoformat() if export.data_as_of else None, "counts": counts, "total_counts": total_counts, "detail_limit": EXPORT_DETAIL_LIMIT, "truncated_sections": truncated_sections, "approval_before_final": True, "source_files": file_rows, "evidence": evidence_rows}
        manifest_path = project_dir / "evidence_manifest.json"
        manifest_path.write_text(_json(manifest), encoding="utf-8")
        _write_csv(project_dir / "warnings.csv", list(warning_rows[0].keys()) if warning_rows else ["id", "warning_type", "severity", "title", "detail", "status", "rule_code", "created_at"], warning_rows)
        _write_csv(project_dir / "drawings.csv", list(drawing_rows[0].keys()) if drawing_rows else ["id", "drawing_number", "change_type", "location_ref", "confidence", "status", "baseline_revision", "changed_revision", "sheet_number"], drawing_rows)
        _write_csv(project_dir / "prices.csv", list(price_rows[0].keys()) if price_rows else ["id", "candidate_id", "item_name", "specification", "unit", "price", "lookup_status", "service_name", "source_file_id", "reference_date"], price_rows)
        _write_csv(project_dir / "approvals.csv", list(approval_rows[0].keys()) if approval_rows else ["id", "source_id", "department", "decision", "reviewer", "comment", "evidence_ref", "created_at"], approval_rows)
        _write_csv(project_dir / "evidence.csv", list(evidence_rows[0].keys()) if evidence_rows else ["id", "warning_id", "source_file_id", "file_path", "sheet_name", "row_ref", "cell_ref", "page_number", "location_text", "extraction_confidence", "evidence_note"], evidence_rows)
        _write_csv(project_dir / "source_files.csv", list(file_rows[0].keys()) if file_rows else ["id", "original_name", "file_path", "sha256", "document_type", "revision", "drawing_number", "sheet_name", "is_valid"], file_rows)
        sheets = {
            "요약": (["항목", "값", "표시 기준"], [{"항목": "프로젝트", "값": export.project_id, "표시 기준": "검토 스냅샷"}, {"항목": "경고", "값": total_counts["warnings"], "표시 기준": f"승인 전 후보 · 상세 {len(warning_rows)}건 포함"}, {"항목": "도면 변경 후보", "값": total_counts["drawings"], "표시 기준": f"근거 확인 대기 · 상세 {len(drawing_rows)}건 포함"}, {"항목": "단가 후보", "값": total_counts["prices"], "표시 기준": f"구매부서 판단 대기 · 상세 {len(price_rows)}건 포함"}, {"항목": "검토 대기", "값": pending, "표시 기준": "자동 확정 금지"}, {"항목": "상세 제한", "값": EXPORT_DETAIL_LIMIT, "표시 기준": "최신순 상세 행 상한 · Manifest에서 생략 여부 확인"}]),
            "도면변경": (list(drawing_rows[0].keys()) if drawing_rows else ["id"], drawing_rows),
            "수량산식": (["warning_id", "title", "detail", "status", "rule_code"], [{"warning_id": row["id"], "title": row["title"], "detail": row["detail"], "status": row["status"], "rule_code": row["rule_code"]} for row in warning_rows if "quantity" in row["warning_type"] or "formula" in row["warning_type"]]),
            "내역연결": (["warning_id", "title", "detail", "status"], [{"warning_id": row["id"], "title": row["title"], "detail": row["detail"], "status": row["status"]} for row in warning_rows if "mapping" in row["warning_type"] or "mismatch" in row["warning_type"]]),
            "단가검토": (list(price_rows[0].keys()) if price_rows else ["id"], price_rows),
            "승인이력": (list(approval_rows[0].keys()) if approval_rows else ["id"], approval_rows),
            "원본근거": (list(evidence_rows[0].keys()) if evidence_rows else ["id"], evidence_rows),
            "데이터정의": (["필드", "설명"], [{"필드": "price", "설명": "승인 전 후보 단가이며 확정값이 아님"}, {"필드": "source_file_id", "설명": "원본 파일 추적 ID"}, {"필드": "export_id", "설명": "동일 검토 시점 스냅샷 ID"}]),
        }
        xlsx_path = _build_xlsx(project_dir / "review_result.xlsx", sheets)
        # PDF에는 내부 ID 대신 검토자가 선택한 표시명을 보여준다. Manifest에는
        # 재현 가능한 원본 ID를 그대로 남긴다.
        pdf_filters = dict(filters)
        if pdf_filters.get("review_stage"):
            pdf_filters["review_stage"] = {"initial": "최초자료 검토", "change": "설계변경 검토", "price": "신규내역 단가"}.get(pdf_filters["review_stage"], pdf_filters["review_stage"])
        if pdf_filters.get("building_id"):
            building = db.get(Building, pdf_filters["building_id"])
            pdf_filters["building_id"] = building.name if building else pdf_filters["building_id"]
        if pdf_filters.get("work_package_id"):
            work_package = db.get(WorkPackage, pdf_filters["work_package_id"])
            pdf_filters["work_package_id"] = work_package.name if work_package else pdf_filters["work_package_id"]
        pdf_path = _build_pdf(project_dir / "summary.pdf", export.project_id, total_counts, pending, pdf_filters, warning_rows, drawing_rows, price_rows)
        bundle_path = project_dir / "review_result_bundle.zip"
        with zipfile.ZipFile(bundle_path, "w", zipfile.ZIP_DEFLATED) as bundle:
            for child in project_dir.iterdir():
                if child.name != bundle_path.name and child.is_file():
                    bundle.write(child, child.name)
        export.bundle_path = str(bundle_path)
        export.xlsx_path = xlsx_path
        export.pdf_path = pdf_path
        export.manifest_path = str(manifest_path)
        export.status = "completed" if xlsx_path and pdf_path.endswith(".pdf") else "completed_with_warning"
        if not xlsx_path or not pdf_path.endswith(".pdf"):
            export.error_message = "선택적 Excel/PDF 라이브러리가 없어 일부 출력이 CSV 또는 HTML로 대체되었습니다. 운영 이미지에서는 requirements.txt 의 패키지를 설치하세요."
        db.commit()
    except Exception as error:
        db.rollback()
        export = db.get(ReviewExport, export_id)
        if export:
            export.status = "failed"
            export.error_message = str(error)[:2000]
            db.commit()
    finally:
        db.close()
